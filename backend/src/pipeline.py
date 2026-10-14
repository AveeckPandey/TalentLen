"""Batch processing pipeline with bounded concurrency, error handling, and reporting."""

import csv
import json
import logging
from pathlib import Path
from typing import List, Optional, Tuple
from concurrent.futures import ThreadPoolExecutor, as_completed

from src.config import settings
from src.models import (
    ParsedResume,
    EligibilityResult,
    CandidateResult,
    BatchSummary,
    ScreeningReport,
    FilterPreferences,
    ScoreBreakdown,
    GitHubMetrics,
)
from src.parsers.resume_parser import ResumeParser
from src.filters.eligibility import EligibilityFilter
from src.enrichment.github_client import GitHubClient
from src.scoring.scorer import CandidateScorer
from src.llm.adapter import get_llm_adapter

logger = logging.getLogger(__name__)


class ScreeningPipeline:
    """Orchestrates candidate resume ingestion, filtering, enrichment, and ranking."""

    def __init__(self, github_token: Optional[str] = None):
        self.github_client = GitHubClient(token=github_token)
        self.llm_adapter = get_llm_adapter()

    def process_directory(
        self,
        input_dir: str | Path,
        preferences: Optional[FilterPreferences] = None,
        max_workers: int = 5,
    ) -> ScreeningReport:
        """Process all resumes in a directory concurrently with bounded worker threads."""
        dir_path = Path(input_dir)
        if not dir_path.exists():
            raise FileNotFoundError(f"Input directory does not exist: {dir_path}")

        # Collect supported files
        supported_exts = {".pdf", ".docx", ".doc", ".txt", ".md"}
        resume_files = [
            f for f in dir_path.iterdir()
            if f.is_file() and f.suffix.lower() in supported_exts
        ]
        resume_files.sort(key=lambda p: p.name)

        if not resume_files:
            logger.warning(f"No resume files found in {dir_path}")
            return ScreeningReport(
                batch_summary=BatchSummary(),
                ranked_candidates=[],
                rejected_candidates=[],
            )

        # Process resumes with bounded concurrency
        eligible_candidates: List[CandidateResult] = []
        rejected_candidates: List[CandidateResult] = []
        failed_count = 0
        parsed_count = 0
        github_cohorts_count = {
            "Veteran Builder": 0,
            "Active Builder": 0,
            "Emerging Developer": 0,
            "Dormant / Minimal": 0,
            "Unlinked / Missing": 0,
        }

        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_file = {
                executor.submit(self._process_single_resume, file_path, preferences): file_path
                for file_path in resume_files
            }

            for future in as_completed(future_to_file):
                file_path = future_to_file[future]
                try:
                    cand_result, is_failed = future.result()
                    if is_failed:
                        failed_count += 1
                        rejected_candidates.append(cand_result)
                    else:
                        parsed_count += 1
                        if cand_result.eligible:
                            eligible_candidates.append(cand_result)
                        else:
                            rejected_candidates.append(cand_result)

                        # Tally cohort
                        if cand_result.github_metrics:
                            cohort = cand_result.github_metrics.cohort
                            github_cohorts_count[cohort] = github_cohorts_count.get(cohort, 0) + 1
                        else:
                            github_cohorts_count["Unlinked / Missing"] += 1

                except Exception as exc:
                    logger.error(f"Unexpected error processing {file_path.name}: {exc}")
                    failed_count += 1
                    rejected_candidates.append(
                        CandidateResult(
                            candidate_name=file_path.stem.replace("_", " ").title(),
                            file_name=file_path.name,
                            eligible=False,
                            rejection_reasons=[f"Internal processing error: {str(exc)[:60]}"],
                            reason_50_to_100_words=(
                                f"Candidate {file_path.name} could not be processed due to an unhandled system "
                                f"exception during document ingestion. The file may be corrupt or encoded in an "
                                f"incompatible format. Please verify the resume document integrity."
                            ),
                        )
                    )

        # Sort eligible candidates by total_score descending
        eligible_candidates.sort(key=lambda c: c.total_score, reverse=True)

        # Assign ranks 1 to N
        for idx, candidate in enumerate(eligible_candidates, start=1):
            candidate.rank = idx

        # Sort rejected candidates alphabetically by name
        rejected_candidates.sort(key=lambda c: c.candidate_name)

        # Compute average eligible score
        avg_score = (
            round(sum(c.total_score for c in eligible_candidates) / len(eligible_candidates), 1)
            if eligible_candidates
            else 0.0
        )

        batch_summary = BatchSummary(
            total_resumes=len(resume_files),
            successfully_parsed=parsed_count,
            eligible=len(eligible_candidates),
            rejected=len(rejected_candidates),
            failed_unreadable=failed_count,
            github_cohorts=github_cohorts_count,
            average_score_eligible=avg_score,
        )

        return ScreeningReport(
            batch_summary=batch_summary,
            ranked_candidates=eligible_candidates,
            rejected_candidates=rejected_candidates,
        )

    def _process_single_resume(
        self, file_path: Path, preferences: Optional[FilterPreferences] = None
    ) -> Tuple[CandidateResult, bool]:
        """Process an individual candidate resume through parsing, filtering, scoring, and justification."""
        # 1. Parse File
        resume = ResumeParser.parse_file(file_path)
        if resume.parsing_status != "success":
            failed_result = CandidateResult(
                candidate_name=resume.candidate_name,
                file_name=resume.file_name,
                eligible=False,
                rejection_reasons=[f"Unreadable file: {resume.parsing_status}"],
                matched_skills=[],
                reason_50_to_100_words=(
                    f"The resume file {resume.file_name} could not be parsed successfully ({resume.parsing_status}). "
                    f"Our document ingestion pipeline requires readable text formatting. Please inspect "
                    f"the uploaded file for corruption, password protections, or non-standard font encodings."
                ),
            )
            return failed_result, True

        # 2. Hard Eligibility Check
        eligibility = EligibilityFilter.evaluate(resume, preferences)

        if not eligibility.eligible:
            # Candidate rejected by hard filter
            # Generate narrative reason
            narrative = self.llm_adapter.generate_narrative_reason(
                candidate_name=resume.candidate_name,
                eligible=False,
                total_score=0.0,
                breakdown_dict={},
                strengths=[],
                concerns=[],
                rejection_reasons=eligibility.rejection_reasons,
                matched_skills=eligibility.matched_skills,
                github_summary="Skipped for ineligible candidate",
            )

            rejected_result = CandidateResult(
                rank=None,
                candidate_name=resume.candidate_name,
                file_name=resume.file_name,
                email=resume.email,
                phone=resume.phone,
                resume_url=f"/resumes/{resume.file_name}",
                eligible=False,
                total_score=0.0,
                score_breakdown=ScoreBreakdown(),
                matched_skills=eligibility.matched_skills,
                project_summary="Profile did not meet minimum Python and AI/agentic requirements.",
                github_summary="Not enriched (disqualified before scoring).",
                strengths=[],
                concerns=["Failed hard eligibility criteria"],
                rejection_reasons=eligibility.rejection_reasons,
                reason_50_to_100_words=narrative,
                cgpa=resume.cgpa,
                has_related_experience=resume.has_related_experience,
            )
            return rejected_result, False

        # 3. Eligible Candidate: Enrich with GitHub Activity
        gh_metrics, gh_score, gh_summary = self.github_client.enrich_profile(resume.github_username)

        # 4. Score Candidate (100-pt rubric)
        total_score, breakdown, strengths, concerns, penalties, project_summary = (
            CandidateScorer.score_candidate(resume, gh_metrics, gh_score)
        )

        # 5. Generate 50-100 Word Narrative Justification
        narrative = self.llm_adapter.generate_narrative_reason(
            candidate_name=resume.candidate_name,
            eligible=True,
            total_score=total_score,
            breakdown_dict=breakdown.model_dump(),
            strengths=strengths,
            concerns=concerns,
            rejection_reasons=[],
            matched_skills=eligibility.matched_skills,
            github_summary=gh_summary,
        )

        result = CandidateResult(
            rank=None,  # Assigned after batch sorting
            candidate_name=resume.candidate_name,
            file_name=resume.file_name,
            email=resume.email,
            phone=resume.phone,
            resume_url=f"/resumes/{resume.file_name}",
            eligible=True,
            total_score=total_score,
            score_breakdown=breakdown,
            matched_skills=eligibility.matched_skills,
            project_summary=project_summary,
            github_summary=gh_summary,
            strengths=strengths,
            concerns=concerns,
            rejection_reasons=[],
            reason_50_to_100_words=narrative,
            github_metrics=gh_metrics,
            cgpa=resume.cgpa,
            has_related_experience=resume.has_related_experience,
            penalties_applied=penalties,
        )

        return result, False

    @staticmethod
    def export_json(report: ScreeningReport, output_path: str | Path) -> None:
        """Export screening report to structured JSON file."""
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(report.model_dump(), f, indent=2, ensure_ascii=False)

    @staticmethod
    def export_csv(report: ScreeningReport, output_path: str | Path) -> None:
        """Export screening results to tabular CSV file."""
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        
        all_candidates = report.ranked_candidates + report.rejected_candidates
        fieldnames = [
            "rank", "candidate_name", "email", "phone", "file_name", "eligible", "total_score",
            "ai_project_depth", "python_backend", "cloud_fullstack", "github", "engineering_depth",
            "cgpa", "has_related_experience", "github_cohort", "github_years",
            "github_repos", "github_original_repos", "github_contributions",
            "github_commits", "github_prs", "github_stars",
            "matched_skills", "rejection_reasons", "reason_narrative"
        ]

        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for c in all_candidates:
                writer.writerow({
                    "rank": c.rank if c.rank is not None else "Rejected",
                    "candidate_name": c.candidate_name,
                    "email": c.email or "N/A",
                    "phone": c.phone or "N/A",
                    "file_name": c.file_name,
                    "eligible": c.eligible,
                    "total_score": c.total_score,
                    "ai_project_depth": c.score_breakdown.ai_project_depth,
                    "python_backend": c.score_breakdown.python_backend,
                    "cloud_fullstack": c.score_breakdown.cloud_fullstack,
                    "github": c.score_breakdown.github,
                    "engineering_depth": c.score_breakdown.engineering_depth,
                    "cgpa": c.cgpa if c.cgpa is not None else "N/A",
                    "has_related_experience": c.has_related_experience,
                    "github_cohort": c.github_metrics.cohort if c.github_metrics else "N/A",
                    "github_years": c.github_metrics.account_age_years if c.github_metrics else 0.0,
                    "github_repos": c.github_metrics.public_repos if c.github_metrics else 0,
                    "github_original_repos": c.github_metrics.original_repos_count if c.github_metrics else 0,
                    "github_contributions": c.github_metrics.total_contributions if c.github_metrics else 0,
                    "github_commits": c.github_metrics.recent_commits if c.github_metrics else 0,
                    "github_prs": c.github_metrics.prs_count if c.github_metrics else 0,
                    "github_stars": c.github_metrics.total_stars if c.github_metrics else 0,
                    "matched_skills": "; ".join(c.matched_skills),
                    "rejection_reasons": "; ".join(c.rejection_reasons),
                    "reason_narrative": c.reason_50_to_100_words,
                })

    @staticmethod
    def export_excel(report: ScreeningReport, output_path: str | Path) -> None:
        """Export screening results to a styled Excel (.xlsx) spreadsheet using openpyxl."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Candidate Shortlist"

        headers = [
            "Rank", "Candidate Name", "Email", "Phone", "Resume File", "Status",
            "Total Score", "AI Depth (40)", "Python (30)", "Cloud (15)", "GitHub (10)", "Eng (5)",
            "CGPA", "Experience", "GitHub Cohort", "GitHub Years", "Public Repos", "Original Repos",
            "Total Contributions", "Commits", "PRs", "Stars", "Matched Skills", "Evaluation Justification"
        ]

        # Header styling
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if isinstance(report, list):
            all_candidates = report
        else:
            all_candidates = report.ranked_candidates + report.rejected_candidates
        for row_idx, c in enumerate(all_candidates, start=2):
            row_data = [
                c.rank if c.rank is not None else "Disqualified",
                c.candidate_name,
                c.email or "N/A",
                c.phone or "N/A",
                c.file_name,
                "Eligible" if c.eligible else "Disqualified",
                c.total_score,
                c.score_breakdown.ai_project_depth,
                c.score_breakdown.python_backend,
                c.score_breakdown.cloud_fullstack,
                c.score_breakdown.github,
                c.score_breakdown.engineering_depth,
                c.cgpa if c.cgpa is not None else "N/A",
                "Yes" if c.has_related_experience else "No",
                c.github_metrics.cohort if c.github_metrics else "N/A",
                c.github_metrics.account_age_years if c.github_metrics else 0.0,
                c.github_metrics.public_repos if c.github_metrics else 0,
                c.github_metrics.original_repos_count if c.github_metrics else 0,
                c.github_metrics.total_contributions if c.github_metrics else 0,
                c.github_metrics.recent_commits if c.github_metrics else 0,
                c.github_metrics.prs_count if c.github_metrics else 0,
                c.github_metrics.total_stars if c.github_metrics else 0,
                ", ".join(c.matched_skills),
                c.reason_50_to_100_words,
            ]
            ws.append(row_data)

            # Apply cell formatting
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                if col_idx in [1, 6, 7, 13, 14]:
                    cell.alignment = Alignment(horizontal="center")
                if not c.eligible and col_idx == 6:
                    cell.font = Font(color="9C0006", bold=True)
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif c.eligible and col_idx == 6:
                    cell.font = Font(color="006100", bold=True)
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        # Set column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(40, max(12, max_len + 2))

        wb.save(path)

    @staticmethod
    def export_docx(report: ScreeningReport | List[CandidateResult], output_path: str | Path) -> None:
        """Export screening results to a formal Word document (.docx) using python-docx."""
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # Title
        title = doc.add_heading("Candidate Resume Screening & Ranking Report", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = doc.add_paragraph("Role: SDE Internship (Python Fundamentals & Agentic AI)")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Summary Section
        doc.add_heading("1. Summary of Evaluated Candidates", level=1)
        if isinstance(report, list):
            ranked_candidates = [c for c in report if c.eligible] or report
            bs_total = len(report)
            bs_eligible = len([c for c in report if c.eligible])
            bs_rejected = len([c for c in report if not c.eligible])
            avg_score = round(sum(c.total_score for c in ranked_candidates) / len(ranked_candidates), 1) if ranked_candidates else 0.0
        else:
            ranked_candidates = report.ranked_candidates
            bs = report.batch_summary
            bs_total = bs.total_resumes
            bs_eligible = bs.eligible
            bs_rejected = bs.rejected
            avg_score = bs.average_score_eligible

        p_summary = doc.add_paragraph()
        p_summary.add_run(f"• Total Candidates Selected: {bs_total}\n")
        p_summary.add_run(f"• Eligible Candidates: {bs_eligible}\n")
        p_summary.add_run(f"• Disqualified Candidates: {bs_rejected}\n")
        p_summary.add_run(f"• Average Score (Eligible): {avg_score} / 100\n")

        # Top Shortlisted Candidates Table
        doc.add_heading("2. Shortlisted Candidates", level=1)
        table = doc.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Rank"
        hdr_cells[1].text = "Candidate Name"
        hdr_cells[2].text = "Email"
        hdr_cells[3].text = "Phone"
        hdr_cells[4].text = "Score"
        hdr_cells[5].text = "GitHub Cohort"

        for c in ranked_candidates:
            row_cells = table.add_row().cells
            row_cells[0].text = f"#{c.rank}" if c.rank is not None else "Disq."
            row_cells[1].text = c.candidate_name
            row_cells[2].text = c.email or "N/A"
            row_cells[3].text = c.phone or "N/A"
            row_cells[4].text = f"{c.total_score:.1f}"
            row_cells[5].text = c.github_metrics.cohort if c.github_metrics else "N/A"

        # Detailed Profiles Section
        doc.add_heading("3. Detailed Candidate Evaluations", level=1)
        for c in ranked_candidates[:20]:
            doc.add_heading(f"#{c.rank or 'Disq.'} - {c.candidate_name} ({c.total_score:.1f}/100)", level=2)
            p = doc.add_paragraph()
            p.add_run(f"Contact: {c.email or 'N/A'} | {c.phone or 'N/A'} | Resume: {c.file_name}\n").bold = True
            p.add_run(f"Score Breakdown: AI Depth: {c.score_breakdown.ai_project_depth} | Python: {c.score_breakdown.python_backend} | Cloud: {c.score_breakdown.cloud_fullstack} | GitHub: {c.score_breakdown.github} | Eng: {c.score_breakdown.engineering_depth}\n")
            if c.github_metrics:
                p.add_run(f"GitHub: {c.github_metrics.cohort} ({c.github_metrics.account_age_years} yrs, {c.github_metrics.public_repos} repos, {c.github_metrics.total_contributions} contributions, {c.github_metrics.prs_count} PRs)\n")
            p.add_run(f"Evaluation Notes: {c.reason_50_to_100_words}\n")

        doc.save(path)
